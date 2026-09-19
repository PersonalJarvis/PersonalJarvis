"""Material: an order text becomes a slide deck (PPTX + PDF) with a quality check.

Flow: the active brain writes a JSON outline -> ``python-pptx`` renders the
PPTX -> a Chromium-family browser already on the machine (Edge on Windows)
prints the same slides to PDF -> deterministic checks. Everything lands in
``<user_data_dir>/materials/<stamp>/``; the HTML used for the PDF lives in a
temporary directory that is removed afterwards. Nothing is sent anywhere:
delivering the files is the user's decision.

``python-pptx`` is the optional ``[material]`` extra; without it the PDF is
still produced and the missing PPTX is reported, not faked.
"""

from __future__ import annotations

import html
import json
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from jarvis.core.paths import user_data_dir
from jarvis.core.process_utils import NO_WINDOW_CREATIONFLAGS

MAX_BULLETS = 6
MAX_BULLET_CHARS_CJK = 60
MAX_BULLET_CHARS_LATIN = 110
_COUNT_RE = re.compile(
    r"(\d{1,2})\s*(\u679a|\u30da\u30fc\u30b8|\u30b9\u30e9\u30a4\u30c9|slides?|pages?)", re.I
)
_CJK_RE = re.compile(r"[\u3040-\u30ff\u4e00-\u9fff]")

OUTLINE_SYSTEM = (
    "You design presentation decks. Reply with JSON only, no prose, in this "
    'shape: {{"title": str, "subtitle": str, "slides": [{{"title": str, '
    '"bullets": [str, ...]}}]}}. {count}Use at most 6 short bullets per slide '
    "(one line each). Write every text in {language}. Do not invent figures or "
    "facts that the order does not give; say what the slide should contain instead."
)


@dataclass
class Deck:
    title: str
    subtitle: str = ""
    slides: list[dict[str, Any]] = field(default_factory=list)


def requested_slide_count(order: str) -> int | None:
    m = _COUNT_RE.search(order or "")
    return int(m.group(1)) if m else None


def outline_system(order: str, language: str) -> str:
    n = requested_slide_count(order)
    count = f"Make exactly {n} content slides. " if n else "Make 5 to 10 content slides. "
    return OUTLINE_SYSTEM.format(count=count, language=language)


def parse_outline(text: str) -> Deck:
    """The first JSON object in ``text`` as a Deck; raises ValueError if none."""
    start, end = (text or "").find("{"), (text or "").rfind("}")
    if start < 0 or end <= start:
        raise ValueError("the model returned no JSON outline")
    data = json.loads(text[start : end + 1])
    slides = []
    for s in data.get("slides") or []:
        if not isinstance(s, dict):
            continue
        bullets = [str(b).strip() for b in (s.get("bullets") or []) if str(b).strip()]
        slides.append({"title": str(s.get("title") or "").strip(), "bullets": bullets})
    return Deck(
        title=str(data.get("title") or "").strip(),
        subtitle=str(data.get("subtitle") or "").strip(),
        slides=slides,
    )


def output_dir(stamp: str | None = None) -> Path:
    path = user_data_dir() / "materials" / (stamp or datetime.now().strftime("%Y%m%d-%H%M%S"))
    path.mkdir(parents=True, exist_ok=True)
    return path


def render_pptx(deck: Deck, path: Path) -> str | None:
    """Write the PPTX; returns an error string instead of raising."""
    try:
        from pptx import Presentation  # noqa: PLC0415 - optional [material] extra
        from pptx.util import Inches, Pt  # noqa: PLC0415
    except ImportError:
        return "python-pptx is not installed (pip install python-pptx)"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    first = prs.slides.add_slide(prs.slide_layouts[0])
    first.shapes.title.text = deck.title
    if len(first.placeholders) > 1:
        first.placeholders[1].text = deck.subtitle
    for s in deck.slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = s["title"]
        body = slide.placeholders[1].text_frame
        body.clear()
        for i, bullet in enumerate(s["bullets"]):
            para = body.paragraphs[0] if i == 0 else body.add_paragraph()
            para.text = bullet
            para.font.size = Pt(24)
    prs.save(str(path))
    return None


def _browser() -> str | None:
    for name in ("msedge", "chrome", "chromium", "chromium-browser", "google-chrome"):
        found = shutil.which(name)
        if found:
            return found
    if sys.platform == "win32":
        for base in (r"C:\Program Files (x86)", r"C:\Program Files"):
            for rel in (
                r"Microsoft\Edge\Application\msedge.exe",
                r"Google\Chrome\Application\chrome.exe",
            ):
                candidate = Path(base) / rel
                if candidate.is_file():
                    return str(candidate)
    return None


def deck_html(deck: Deck) -> str:
    e = html.escape
    pages = [f'<section class="t"><h1>{e(deck.title)}</h1><p>{e(deck.subtitle)}</p></section>']
    for s in deck.slides:
        items = "".join(f"<li>{e(b)}</li>" for b in s["bullets"])
        pages.append(f"<section><h2>{e(s['title'])}</h2><ul>{items}</ul></section>")
    return (
        "<!doctype html><meta charset=utf-8><style>"
        "@page{size:13.333in 7.5in;margin:0}body{margin:0;font-family:'Yu Gothic','Meiryo',"
        "'Noto Sans CJK JP',sans-serif}section{width:13.333in;height:7.5in;padding:.7in;"
        "box-sizing:border-box;page-break-after:always}h1{font-size:44pt;margin-top:2in}"
        "h2{font-size:34pt;border-bottom:3px solid #333}li{font-size:22pt;margin:.15in 0}"
        "</style>" + "".join(pages)
    )


def render_pdf(deck: Deck, path: Path) -> str | None:
    """Print the deck to PDF with a local Chromium browser; error string on failure."""
    browser = _browser()
    if browser is None:
        return "no Edge/Chrome found to print the PDF"
    with tempfile.TemporaryDirectory(prefix="jarvis-material-") as tmp:
        page = Path(tmp) / "deck.html"
        page.write_text(deck_html(deck), encoding="utf-8")
        out = Path(tmp) / "deck.pdf"
        cmd = [
            browser,
            "--headless=new",
            "--disable-gpu",
            "--no-pdf-header-footer",
            f"--user-data-dir={Path(tmp) / 'profile'}",
            f"--print-to-pdf={out}",
            page.as_uri(),
        ]
        try:
            subprocess.run(
                cmd,
                capture_output=True,
                timeout=90,
                check=False,
                creationflags=NO_WINDOW_CREATIONFLAGS,
            )
        except (OSError, subprocess.TimeoutExpired) as exc:
            return f"PDF printing failed: {exc}"
        if not out.is_file() or out.stat().st_size == 0:
            return "the browser produced no PDF"
        shutil.move(str(out), path)
    return None


@dataclass
class Check:
    name: str
    ok: bool
    detail: str = ""


def quality_check(deck: Deck, order: str, pptx: Path | None, pdf: Path | None) -> list[Check]:
    checks: list[Check] = []
    wanted = requested_slide_count(order)
    n = len(deck.slides)
    checks.append(
        Check(
            "slide_count",
            wanted is None or n == wanted,
            f"{n} slides" + (f" (ordered {wanted})" if wanted else ""),
        )
    )
    checks.append(Check("deck_title", bool(deck.title)))
    untitled = [i + 1 for i, s in enumerate(deck.slides) if not s["title"]]
    checks.append(Check("slide_titles", not untitled, f"missing on {untitled}" if untitled else ""))
    empty = [i + 1 for i, s in enumerate(deck.slides) if not s["bullets"]]
    checks.append(Check("no_empty_slides", not empty, f"empty: {empty}" if empty else ""))
    crowded = [i + 1 for i, s in enumerate(deck.slides) if len(s["bullets"]) > MAX_BULLETS]
    checks.append(
        Check(
            "bullets_per_slide",
            not crowded,
            f"more than {MAX_BULLETS} on {crowded}" if crowded else "",
        )
    )
    long_ = []
    for i, s in enumerate(deck.slides):
        for b in s["bullets"]:
            limit = MAX_BULLET_CHARS_CJK if _CJK_RE.search(b) else MAX_BULLET_CHARS_LATIN
            if len(b) > limit:
                long_.append(i + 1)
                break
    checks.append(Check("text_length", not long_, f"long text on {long_}" if long_ else ""))
    titles = [s["title"] for s in deck.slides if s["title"]]
    dup = sorted({t for t in titles if titles.count(t) > 1})
    checks.append(Check("unique_titles", not dup, ", ".join(dup)))
    if pptx is not None:
        try:
            from pptx import Presentation  # noqa: PLC0415

            count = len(Presentation(str(pptx)).slides)
            checks.append(Check("pptx_opens", count == n + 1, f"{count} pages"))
        except Exception as exc:  # noqa: BLE001 - a broken file is a failed check
            checks.append(Check("pptx_opens", False, str(exc)[:120]))
    checks.append(
        Check("pdf_written", pdf is not None and pdf.is_file() and pdf.stat().st_size > 0)
    )
    return checks


def report(checks: list[Check]) -> str:
    lines = [
        f"{'OK ' if c.ok else 'NG '} {c.name}" + (f": {c.detail}" if c.detail else "")
        for c in checks
    ]
    return "\n".join(lines)


__all__ = [
    "Check",
    "Deck",
    "deck_html",
    "outline_system",
    "output_dir",
    "parse_outline",
    "quality_check",
    "render_pdf",
    "render_pptx",
    "report",
    "requested_slide_count",
]
