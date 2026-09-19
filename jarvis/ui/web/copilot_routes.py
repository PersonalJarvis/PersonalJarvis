"""Copilot dashboard API (``/api/copilot/*``): Material, Teacher, Workflow.

Thin: reads what the copilot features already wrote under the user data dir
and forwards commands to the same BrainManager handlers the chat uses, so a
button and a spoken command do exactly the same thing. Nothing here sends
anything outside the machine.
"""

from __future__ import annotations

import json
import logging
import os
import sys
from pathlib import Path
from typing import Any

from fastapi import APIRouter, HTTPException, Request
from pydantic import BaseModel, Field

from jarvis.core.paths import user_data_dir

log = logging.getLogger(__name__)

router = APIRouter(prefix="/api/copilot", tags=["copilot"])

_AREAS = ("materials", "lessons", "workflow")
_MAX_ITEMS = 20


def _area(name: str) -> Path:
    return user_data_dir() / name


def _read_pptx(path: Path) -> tuple[str, int]:
    """(title, content slide count) of a deck made before deck.json existed."""
    try:
        from pptx import Presentation  # noqa: PLC0415 - optional [material] extra

        prs = Presentation(str(path))
    except Exception:  # noqa: BLE001 - no python-pptx or no readable file: unknown
        log.debug("copilot: %s not readable as PPTX", path)
        return "", 0
    slides = list(prs.slides)
    title = ""
    if slides and slides[0].shapes.title is not None:
        title = slides[0].shapes.title.text
    return title, max(len(slides) - 1, 0)


def _materials() -> list[dict[str, Any]]:
    root = _area("materials")
    if not root.is_dir():
        return []
    out = []
    for folder in sorted((p for p in root.iterdir() if p.is_dir()), reverse=True)[:_MAX_ITEMS]:
        deck: dict[str, Any] = {}
        try:
            deck = json.loads((folder / "deck.json").read_text(encoding="utf-8"))
        except (OSError, ValueError):
            log.debug("copilot: %s has no readable deck.json", folder.name)
        quality = ""
        try:
            quality = (folder / "quality.txt").read_text(encoding="utf-8")
        except OSError:
            log.debug("copilot: %s has no quality.txt", folder.name)
        lines = [ln for ln in quality.splitlines() if ln[:2] in ("OK", "NG")]
        title, slides = deck.get("title", ""), len(deck.get("slides") or [])
        if not deck:
            title, slides = _read_pptx(folder / "deck.pptx")
        out.append(
            {
                "id": folder.name,
                "title": title,
                "order": deck.get("order", ""),
                "slides": slides,
                "passed": sum(ln.startswith("OK") for ln in lines),
                "total": len(lines),
                "quality": quality,
                "files": [p.name for p in sorted(folder.glob("deck.*")) if p.suffix != ".json"],
            }
        )
    return out


def _files(area: str, pattern: str) -> list[dict[str, Any]]:
    root = _area(area)
    if not root.is_dir():
        return []
    files = sorted(root.glob(pattern), key=lambda p: p.stat().st_mtime, reverse=True)
    return [{"name": p.name, "size": p.stat().st_size} for p in files[:_MAX_ITEMS]]


@router.get("/overview")
def overview(request: Request) -> dict[str, Any]:
    """Everything the dashboard shows, in one read."""
    from jarvis.copilot import workflow

    brain = getattr(request.app.state, "brain", None)
    lesson = getattr(brain, "_lesson", None)
    plan = getattr(brain, "_lesson_plan", None)
    report = ""
    reports = _files("workflow", "*-report.md")
    if reports:
        try:
            report = (_area("workflow") / reports[0]["name"]).read_text(encoding="utf-8")
        except OSError:
            log.debug("copilot: latest workflow report unreadable")
    return {
        "materials": _materials(),
        "lessons": _files("lessons", "*.md"),
        "lesson": None
        if lesson is None
        else {
            "topic": lesson.topic,
            "minutes": lesson.minutes,
            "elapsed": round(lesson.elapsed_min(), 1),
            "remaining": round(lesson.remaining_min(), 1),
            "utterances": len(lesson.utterances),
            "summaries": lesson.summaries,
        },
        "lesson_plan": None if plan is None else {"topic": plan.topic, "minutes": plan.minutes},
        "observing": getattr(brain, "_wf_observer", None) is not None,
        "workflow_report": report,
        "candidates": [
            {"steps": list(c.steps), "repeats": c.repeats, "seconds_each": c.seconds_each}
            for c in workflow.load_candidates()
        ],
        "drafts": _files("workflow", "*.ps1"),
    }


class RunBody(BaseModel):
    text: str = Field(..., min_length=1, max_length=4000)
    language: str = Field(default="en", max_length=8)


@router.post("/run")
async def run(body: RunBody, request: Request) -> dict[str, Any]:
    """Run one copilot or teacher command, exactly as if it had been said."""
    from jarvis.copilot.gate import match_copilot_command
    from jarvis.teacher.gate import match_teacher_command

    brain = getattr(request.app.state, "brain", None)
    if brain is None:
        raise HTTPException(status_code=503, detail="The assistant is not ready yet.")
    # A button has no spoken turn to take the language from: the UI language
    # decides the reply language.
    lang = body.language if body.language in ("de", "en", "es", "ja") else "en"
    lesson = getattr(brain, "_lesson", None)
    teacher = match_teacher_command(body.text, lesson_active=lesson is not None)
    if teacher is not None:
        return {"reply": await brain._handle_teacher_command(teacher, lang)}
    cmd = match_copilot_command(body.text)
    if cmd is not None:
        return {"reply": await brain._handle_copilot_command(cmd, lang)}
    raise HTTPException(status_code=400, detail="Not a copilot command.")


class OpenBody(BaseModel):
    area: str
    path: str = Field(..., min_length=1, max_length=300)


@router.post("/open")
def open_file(body: OpenBody) -> dict[str, Any]:
    """Open a copilot file (or its folder) with the OS default app, local only."""
    if body.area not in _AREAS:
        raise HTTPException(status_code=400, detail="Unknown area.")
    root = _area(body.area).resolve()
    target = (root / body.path).resolve()
    if root not in target.parents and target != root:
        raise HTTPException(status_code=400, detail="Path outside the copilot folders.")
    if not target.exists():
        raise HTTPException(status_code=404, detail="Not found.")
    if sys.platform != "win32":
        raise HTTPException(status_code=501, detail="Opening files is available on Windows.")
    os.startfile(str(target))  # noqa: S606 - a local file the user clicked, inside our folder
    return {"ok": True}


__all__ = ["router"]
